"""Stage 2 Option A: build the district+tehsil observation panel."""
import argparse, collections, csv, difflib, json, pathlib, re, sys
import openpyxl
from read_workbook import read, validate_spec, merged_ranges
from unit_aliases import apply as apply_alias

from table_spec import SPEC
TABLES = [t for t in SPEC if not SPEC[t].get('locality_list')]
REGIONS = {'kp': 'KHYBER PAKHTUNKHWA', 'punjab': 'PUNJAB', 'sindh': 'SINDH',
           'balochistan': 'BALOCHISTAN', 'islamabad': 'ISLAMABAD'}
PMAP = dict(REGIONS, islamabad='ISLAMABAD CAPITAL TERRITORY')


def norm(s):
    s = re.sub(r'\s+', ' ', s.upper())
    s = re.sub(r'\b(TEHSIL|TALUKA|TALUKO|SUB-?TEHSIL|SUB-?DIVISION|TOWN|DISTRICT|CITY|CANTT)\b', '', s)
    return re.sub(r'[^A-Z0-9]', '', re.sub(r'\(.*?\)', '', s))


def write_observations(obs, cols, out):
    """Parquet when DuckDB is available, CSV otherwise. 2.6M rows of CSV is
    400 MB; the same data is roughly twenty times smaller as Parquet."""
    tmp = out / '_observations.csv'
    with open(tmp, 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction='ignore')
        w.writeheader(); w.writerows(obs)
    try:
        import duckdb
    except ImportError:
        tmp.rename(out / 'observations.csv'); return
    # Pin the schema: '13a' is a table name, not a number, and `value` must
    # stay a double even where a column happens to hold only integers.
    types = {c: 'VARCHAR' for c in cols}
    types['value'] = 'DOUBLE'
    types['missing'] = 'BOOLEAN'
    for c in ('src_row', 'src_col'):
        types[c] = 'INTEGER'
    spec = ', '.join(f"'{k}': '{v}'" for k, v in types.items())
    duckdb.sql(f"COPY (SELECT * FROM read_csv('{tmp}', header=true, quote='\"', escape='\"', "
               f"columns={{{spec}}})) TO '{out / 'observations.parquet'}' "
               f"(FORMAT PARQUET, COMPRESSION ZSTD)")
    tmp.unlink()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', required=True)
    ap.add_argument('--crosswalk', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    src, out = pathlib.Path(a.src), pathlib.Path(a.out)
    out.mkdir(parents=True, exist_ok=True)

    cw = collections.defaultdict(dict)
    for row in csv.DictReader(open(a.crosswalk)):
        # The mouza2020 crosswalk is an external input and keeps its own column
        # names; `province_area` is this panel's column, not that file's.
        cw[row['province']].setdefault(norm(row['tehsil']), row)

    # PBS's province files are not always what their name says: table 6 is
    # published under a KP filename but contains all 135 districts, and there
    # is no KP-only table 6. So the district's area is taken from the
    # table 1 roster rather than from the filename, and observations are
    # de-duplicated afterwards.
    roster = {}
    for reg, prov in REGIONS.items():
        f = src / 'xlsx' / f'table_1_{reg}.xlsx'
        if not f.exists():
            continue
        ws = openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0]
        for r in ws.iter_rows(values_only=True):
            v = r[0]
            if isinstance(v, str) and 'DISTRICT' in v.upper():
                # Table 1 carries the ALL-deletion defect, so the roster has to
                # be keyed on the corrected name or the lookup misses and the
                # district falls back to whatever the filename claims.
                name = apply_alias('1', re.sub(r'\s+', ' ', v.strip()))[0]
                roster[name.upper()] = prov

    obs, units, fams, problems = [], {}, {}, []
    for t in TABLES:
        for reg, prov in REGIONS.items():
            f = src / 'xlsx' / f'table_{t}_{reg}.xlsx'
            if not f.exists():
                problems.append(dict(table=t, region=reg, issue='file missing')); continue
            ws = openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            err = validate_spec(rows, t)
            fams[(t, reg)] = err or 'ok'
            if err:
                problems.append(dict(table=t, region=reg, issue=err))
            n = 0
            for o in read(rows, prov, t, merged_ranges(f)):
                if o['district']:
                    o['province_area'] = roster.get(
                        re.sub(r'\s+', ' ', o['district'].strip().upper()), prov)
                o['home_file'] = (o['province_area'] == prov)
                o['source_file'] = f.name
                o['sheet'] = ws.title
                obs.append(o); n += 1
                # Keyed on district as well as name: Punjab has a SAHIWAL
                # TEHSIL in both Sahiwal and Sargodha districts, and they are
                # different places.
                key = (o['province_area'], o['district'], o['unit'])
                units.setdefault(key, dict(province_area=o['province_area'], district=o['district'],
                                           unit=o['unit'], unit_type=o['unit_type'],
                                           tables=set()))['tables'].add(t)
            if n == 0:
                problems.append(dict(table=t, region=reg, issue='no observations'))

    # De-duplicate: where a district appears in both its own area file and
    # a mislabelled multi-area one, keep the row from the file that belongs
    # to it. Sorting puts home-file rows first, so the first seen wins.
    obs.sort(key=lambda o: (not o['home_file'], o['source_file'], o['src_row'], o['src_col']))
    seen, deduped, dropped = set(), [], 0
    for o in obs:
        k = (o['table_id'], o['province_area'], o['district'], o['unit'], o['locality'],
             o['sex'], o['indicator'], o['col_label'], o['src_col'])
        if k in seen:
            dropped += 1
            continue
        seen.add(k)
        deduped.append(o)
    obs = deduped
    if dropped:
        problems.append(dict(table='-', region='-',
                             issue=f'{dropped} duplicate observations dropped '
                                   f'(multi-area source files)'))

    # polygon join, sub-district units only
    keys = {p: list(d) for p, d in cw.items()}
    for u in units.values():
        if u['unit_type'] == 'district':
            u['dd_id'], u['match'] = '', ''
            continue
        d = cw[PMAP[[k for k, v in REGIONS.items() if v == u['province_area']][0]]]
        n = norm(u['unit']); hit, how = d.get(n), 'exact'
        if not hit:
            g = difflib.get_close_matches(n, keys[PMAP[[k for k, v in REGIONS.items()
                                                        if v == u['province_area']][0]]], n=1, cutoff=0.82)
            if g:
                hit, how = d[g[0]], 'fuzzy'
        u['dd_id'] = hit['dd_id'] if hit else ''
        u['match'] = f"{hit['match']}/{how}" if hit else 'unmatched'

    cols = ['province_area', 'table_id', 'district', 'unit', 'unit_source', 'unit_type', 'locality',
            'sex', 'indicator', 'col_label', 'value', 'missing', 'source_file', 'sheet', 'src_row', 'src_col']
    write_observations(obs, cols, out)
    with open(out / 'units.csv', 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=['province_area', 'district', 'unit', 'unit_type',
                                           'dd_id', 'match', 'tables'])
        w.writeheader()
        for u in units.values():
            w.writerow(dict(u, tables=';'.join(sorted(u.pop('tables')))))

    # ---- validation ----
    subs = [u for u in units.values() if u['unit_type'] != 'district']
    dists = [u for u in units.values() if u['unit_type'] == 'district']
    pop = collections.defaultdict(float)
    for o in obs:
        if o['table_id'] == '1' and o['locality'] == 'all' and o['indicator'].startswith('POPULATION-2023') \
           and 'ALL SEXES' in o['indicator'].upper() and isinstance(o['value'], (int, float)):
            pop[(o['province_area'], o['unit'])] = o['value']
    rep = dict(observations=len(obs), units=len(units), districts=len(dists),
               sub_district_units=len(subs),
               dash_cells=sum(1 for o in obs if o['missing']),
               tables=sorted({o['table_id'] for o in obs}),
               spec_checks={f'{t}_{r}': v for (t, r), v in fams.items()},
               matched_units=sum(1 for u in subs if u['match'] != 'unmatched'),
               problems=problems)
    per_table = collections.Counter(o['table_id'] for o in obs)
    rep['observations_per_table'] = dict(sorted(per_table.items(), key=lambda x: TABLES.index(x[0])))
    (out / 'build_report.json').write_text(json.dumps(rep, indent=2, sort_keys=True))

    print(f"observations   {len(obs):,}")
    print(f"units          {len(units):,}  ({len(dists)} districts, {len(subs)} sub-district)")
    print(f"dash cells     {rep['dash_cells']:,} preserved as missing")
    print(f"polygon join   {rep['matched_units']}/{len(subs)} sub-district units")
    print("per table      " + "  ".join(f"T{k}:{v:,}" for k, v in rep['observations_per_table'].items()))
    print(f"spec checks    {dict(collections.Counter(fams.values()))}")
    for p in problems:
        print(f"  PROBLEM {p}")


if __name__ == '__main__':
    main()
