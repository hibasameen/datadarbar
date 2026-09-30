#!/usr/bin/env python3
"""Pakistan weather-related events: immutable downloads -> local Parquet/CSV.

Run with --help. Source records are never automatically merged across providers.
Downloads stay outside git; this module and its tests are version controlled.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import time
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlencode, urljoin, urlparse

import requests
from bs4 import BeautifulSoup

REPO = Path(__file__).resolve().parents[2]
WORKSPACE = REPO.parent
GDACS = "https://www.gdacs.org/gdacsapi/api/Events/geteventlist/search"
UNOSAT = "https://unosat.org"
HAZARDS = {"FL": "flood", "TC": "tropical_cyclone", "DR": "drought", "WF": "wildfire"}
EMDAT_HAZARDS = {"Flood", "Storm", "Drought", "Extreme temperature", "Wildfire",
                 "Glacial lake outburst flood"}


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(data):
    return hashlib.sha256(data).hexdigest()


def norm(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def number(value):
    if value is None or str(value).strip() in ("", "-", "..", "NA", "N/A"):
        return None
    n = float(str(value).replace(",", ""))
    if not math.isfinite(n) or n < 0:
        raise ValueError(f"Invalid nonnegative measurement: {value!r}")
    return n


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")
    temp.replace(path)


class Store:
    """Content-addressed, checksummed snapshots. Refresh appends, never overwrites raw."""
    def __init__(self, root, refresh=False, offline=False):
        self.root = Path(root)
        self.refresh, self.offline = refresh, offline
        self.root.mkdir(parents=True, exist_ok=True)
        self.manifest = self.root / "manifest.jsonl"
        self.index = {}
        if self.manifest.exists():
            for line in self.manifest.read_text().splitlines():
                row = json.loads(line)
                self.index[row["url"]] = row

    def record(self, url, data, source, content_type="", http_status=None):
        sha = digest(data)
        suffix = Path(urlparse(url).path).suffix.lower()
        if suffix not in (".pdf", ".xlsx", ".csv", ".zip", ".json", ".geojson", ".kml"):
            suffix = ".json" if "json" in content_type else ".html"
        path = Path("objects") / source / (sha + suffix)
        dest = self.root / path
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.exists():
            dest.write_bytes(data)
        row = dict(source=source, url=url, retrieved_at=now(), sha256=sha,
                   bytes=len(data), path=str(path), content_type=content_type, http_status=http_status)
        with self.manifest.open("a") as f:
            f.write(json.dumps(row) + "\n")
        self.index[url] = row
        return data, row

    def get(self, url, source):
        old = self.index.get(url)
        if old and (self.offline or not self.refresh):
            data = (self.root / old["path"]).read_bytes()
            if digest(data) != old["sha256"]:
                raise ValueError(f"Checksum mismatch: {old['path']}")
            return data, old
        if self.offline:
            raise FileNotFoundError(f"Not cached: {url}")
        for attempt in range(3):
            try:
                r = requests.get(url, timeout=(15, 45),
                                 headers={"User-Agent": "DataDarbarClimateETL/1.0"})
                if r.status_code == 429 or r.status_code >= 500:
                    r.raise_for_status()
                r.raise_for_status()
                return self.record(url, r.content, source, r.headers.get("Content-Type", ""), r.status_code)
            except requests.RequestException as e:
                if isinstance(e, requests.HTTPError) and e.response is not None:
                    if e.response.status_code < 500 and e.response.status_code != 429:
                        raise
                if attempt == 2:
                    raise
                time.sleep(2 ** attempt)


def gdacs_rows(payload, origin):
    if payload.get("type") != "FeatureCollection" or not isinstance(payload.get("features"), list):
        raise ValueError("GDACS response is not a FeatureCollection")
    rows = []
    for feature in payload["features"]:
        p = feature["properties"]
        countries = {x.get("iso3") for x in p.get("affectedcountries", [])}
        if p.get("iso3") != "PAK" and "PAK" not in countries and p.get("country") != "Pakistan":
            continue
        if p["eventtype"] not in HAZARDS:
            continue
        # A centroid is NOT an affected district. No spatial district assignment here.
        rows.append(dict(record_id=f"gdacs:{p['eventtype']}:{p['eventid']}:PAK",
                         source="gdacs", source_event_id=str(p["eventid"]), country_iso3="PAK",
                         hazard=HAZARDS[p["eventtype"]], subtype=p["eventtype"],
                         title=p.get("name"), start_date=(p.get("fromdate") or "")[:10] or None,
                         end_date=(p.get("todate") or "")[:10] or None, date_precision="day",
                         start_year=int(p["fromdate"][:4]) if p.get("fromdate") else None,
                         glide=p.get("glide"), alert_level=p.get("alertlevel"),
                         deaths=None, affected=None, damage_usd=None,
                         source_url=p.get("url", {}).get("report", origin["url"]),
                         source_modified=p.get("datemodified"), episode_id=str(p.get("episodeid", "")),
                         raw_sha256=origin["sha256"], retrieved_at=origin["retrieved_at"],
                         coverage_note="Pakistan affected; event footprint may cross borders. Alert, not final losses."))
    return rows


def fetch_gdacs(store, start, end):
    all_rows = []
    # Annual windows avoid slow unbounded archive queries. Page 1 is the API's first page.
    for year in range(start.year, end.year + 1):
        lo, hi = max(start, date(year, 1, 1)), min(end, date(year, 12, 31))
        seen = set()
        for page in range(1, 1001):
            params = dict(eventlist="FL;TC;DR;WF", country="Pakistan",
                          fromDate=lo.isoformat(), toDate=hi.isoformat(),
                          pageNumber=page, pageSize=100)
            url = GDACS + "?" + urlencode(params)
            data, meta = store.get(url, "gdacs")
            payload = dict(type="FeatureCollection", features=[]) if meta.get("http_status") == 204 else json.loads(data)
            rows = gdacs_rows(payload, meta)
            fingerprint = digest(json.dumps(payload["features"], sort_keys=True).encode())
            if payload["features"] and fingerprint in seen:
                raise ValueError("GDACS repeated a page; refusing a silently truncated import")
            seen.add(fingerprint)
            all_rows.extend(rows)
            if len(payload["features"]) < 100:
                break
        else:
            raise ValueError("GDACS pagination exceeded 1000 pages")
        print(f"GDACS {year}: {len(all_rows)} Pakistan records collected", flush=True)
    return all_rows


def unosat_assets(product):
    pid = str(product["id"])
    links = []
    for key, kind in [("pdf_name", "pdf"), ("excel_table", "xlsx"),
                      ("shp_link", "shapefile"), ("gdp_link", "geodatabase"), ("kml_link", "kml")]:
        value = product.get(key)
        if not value or str(value).split("/")[-1] == "None":
            continue
        if key == "pdf_name" and "/" not in value:
            value = f"/static/unosat_filesystem/{pid}/{value}"
        elif value.startswith("/unosat_filesystem/"):
            value = "/static" + value
        links.append((kind, urljoin(UNOSAT, value)))
    return links


PROVINCES = {"balochistan", "punjab", "sindh", "khyber pakhtunkhwa", "kpk", "gilgit baltistan",
             "azad kashmir", "azad jammu and kashmir", "fata", "islamabad capital territory",
             "federal capital territory"}


def district_match(name, known):
    # Exact normalized names only. Boundary splits and spelling variants need review.
    key = norm(name)
    return key if key in known else None


def exposure_rows(data, product, meta, known):
    """Verified UNOSAT 3343 layout: one canonical sheet, never sum repeated sheets."""
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet_name = "Statistics the whole country"
    if sheet_name not in workbook.sheetnames:
        raise ValueError("Unknown UNOSAT workbook layout; cached for a new explicit adapter")
    sheet = workbook[sheet_name]
    values = iter(sheet.values)
    header = [norm(x) for x in next(values)]
    expected = ["province district", "area of province district", "analyzed area in cloud free zones km2",
                "percentage of analyzed area", "maximum flood water extent km2",
                "total population in province district", "total population in cloud free area",
                "population potentially exposed", "population potentially exposed"]
    if header[:9] != expected:
        raise ValueError(f"UNOSAT exposure header changed: {header}")
    rows, province = [], None
    for index, r in enumerate(values, 2):
        if not r[0] or not str(r[0]).strip():
            continue
        name = str(r[0]).strip()
        key = norm(name)
        if key == "pakistan":
            level = "country"
        elif key in PROVINCES:
            level, province = "province", name
        elif key == "islamabad":
            level, province = "district", "Islamabad"
        else:
            level = "district"
        # Reject footnotes / unrecognised numeric cells instead of inventing observations.
        area = number(r[1])
        if area is None:
            continue
        dk = district_match(name, known) if level == "district" else None
        rows.append(dict(observation_id=f"unosat:{product['id']}:{sheet_name}:{index}",
                         report_id=f"unosat:{product['id']}", source="unosat",
                         location_name=name, admin_level=level, province=province if level != "country" else None,
                         district_key=dk, match_status="exact_name_boundary_unverified" if dk else "unmatched",
                         observation_start="2022-07-12" if product["id"] == 3343 else None,
                         observation_end="2022-07-21" if product["id"] == 3343 else None,
                         area_km2=area, analyzed_area_km2=number(r[2]),
                         analyzed_fraction=number(r[3]), flooded_area_km2=number(r[4]),
                         population_total=number(r[5]), population_analyzed=number(r[6]),
                         population_exposed=number(r[7]), exposed_fraction=number(r[8]),
                         source_url=meta["url"], source_sheet=sheet_name, source_row=index,
                         raw_sha256=meta["sha256"],
                         measurement_note="Potential satellite/modelled exposure, not reported affected people. "
                         "Country/province rows overlap district rows. '-' retained as missing. "
                         "Name matches do not establish boundary equivalence."))
    workbook.close()
    return rows


def pdf_pages(data, report_id, meta):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return [dict(report_id=report_id, page=i, text=p.extract_text() or "",
                 raw_sha256=meta["sha256"], source_url=meta["url"])
            for i, p in enumerate(reader.pages, 1)]


def fetch_unosat(store, ids, known, download_geometry=False):
    reports, assets, exposures, pages, errors = [], [], [], [], []
    for pid in ids:
        url = f"{UNOSAT}/our_products/{pid}"
        data, meta = store.get(url, "unosat")
        p = json.loads(data)["map_event"]
        if "PAK" not in (p.get("glide") or "") and "pakistan" not in p["title"].lower():
            raise ValueError(f"UNOSAT product {pid} is not identified as Pakistan")
        rid = f"unosat:{pid}"
        reports.append(dict(report_id=rid, source="unosat", title=p["title"],
                            published_at=p.get("created_at"), glide=p.get("glide"),
                            description=p.get("description"), source_notes=p.get("sources"),
                            source_url=f"{UNOSAT}/products/{pid}", raw_sha256=meta["sha256"],
                            retrieved_at=meta["retrieved_at"],
                            record_note="Product publication date is not event/observation date."))
        for kind, link in unosat_assets(p):
            asset = dict(asset_id=digest(link.encode()), report_id=rid, source="unosat",
                         kind=kind, source_url=link, download_status="catalogued", raw_path=None, raw_sha256=None)
            if kind in ("pdf", "xlsx") or download_geometry:
                try:
                    b, m = store.get(link, "unosat")
                    if kind == "pdf" and not b.startswith(b"%PDF"):
                        raise ValueError("Expected PDF bytes")
                    asset.update(download_status="downloaded", raw_path=m["path"], raw_sha256=m["sha256"])
                    if kind == "pdf":
                        pages.extend(pdf_pages(b, rid, m))
                    elif kind == "xlsx":
                        exposures.extend(exposure_rows(b, p, m, known))
                except Exception as exc:
                    asset["download_status"] = "extraction_failed" if asset["raw_path"] else "download_failed"
                    errors.append(dict(source="unosat", url=link, error=str(exc)))
            assets.append(asset)
        print(f"UNOSAT {pid}: {len(exposures)} exposure rows", flush=True)
    return dict(reports=reports, assets=assets, exposures=exposures, pages=pages, errors=errors)


def ndma_links(html, base):
    soup = BeautifulSoup(html, "html.parser")
    links = {}
    for a in soup.find_all("a", href=True):
        url = urljoin(base, a["href"])
        if urlparse(url).path.lower().endswith(".pdf"):
            links[url] = a.get_text(" ", strip=True)
    return sorted(links.items())


def ndma_impacts(pages):
    """Strict adapter for the observed 2026 cumulative provincial table layouts."""
    output = []
    regions = {"Punjab", "KP", "Sindh", "Balochistan", "GB", "AJ&K", "ICT", "Grand Total"}
    for page_index, page in enumerate(pages):
        text = page["text"]
        heading = re.search(r"Cumulative (?:Casualties and Injuries|Damages of Infrastructure & Private Properties)", text)
        if not heading:
            continue
        text = text[heading.start():]
        lines = [(page["page"], line) for line in text.splitlines()]
        # Tables may start on the following PDF page; preserve the actual row page.
        if page_index + 1 < len(pages) and pages[page_index + 1]["report_id"] == page["report_id"]:
            following = pages[page_index + 1]
            lines.extend((following["page"], line) for line in following["text"].splitlines())
        flat = " ".join(" ".join(line for _, line in lines).split())
        if text.startswith("Cumulative Casualties and Injuries"):
            metrics = ["deaths_male", "deaths_female", "deaths_children", "deaths_total",
                       "injured_male", "injured_female", "injured_children", "injured_total"]
            expected_header = "Male Female Children Total Male Female Children Total"
        elif text.startswith("Cumulative Damages of Infrastructure & Private Properties"):
            metrics = ["roads_damaged_km", "bridges_damaged", "houses_destroyed",
                       "houses_partially_damaged", "houses_damaged_total", "livestock_perished"]
            expected_header = "Full Partial Total"
        else:
            continue
        if expected_header not in flat:
            raise ValueError(f"NDMA page {page['page']}: cumulative table header changed")
        period = re.search(r"From (\d{1,2} \w+ \d{4}) to (\d{1,2} \w+ \d{4})", flat)
        if not period:
            raise ValueError("NDMA cumulative table lacks an explicit reporting period")
        start, end = [datetime.strptime(v, "%d %B %Y").date().isoformat() for v in period.groups()]
        observed, row_pages = {}, {}
        for page_number, line in lines:
            line = re.sub(r"Grand T\s*otal", "Grand Total", line.strip())
            # In the 3 September PDF the next chart title touches Punjab's last cell.
            line = line.removesuffix("Cumulative House Damaged").rstrip()
            match = re.fullmatch(r"(Punjab|KP|Sindh|Balochistan|GB|AJ&K|ICT|Grand Total)\s+([\d.,\s-]+)", line)
            if not match:
                continue
            cells = match[2].split()
            if len(cells) != len(metrics) or match[1] in observed:
                raise ValueError("NDMA cumulative table row length/identity changed")
            observed[match[1]] = [number(v) for v in cells]
            row_pages[match[1]] = page_number
            if match[1] == "Grand Total":
                break
        if set(observed) != regions:
            raise ValueError(f"NDMA cumulative table missing regions: {regions - set(observed)}")
        for region, cells in observed.items():
            totals = [(0, 3), (4, 7)] if metrics[0] == "deaths_male" else [(2, 4)]
            for lo, hi in totals:
                group = cells[lo:hi + 1]
                if all(v is not None for v in group) and abs(sum(group[:-1]) - group[-1]) > 0.01:
                    raise ValueError(f"NDMA {region}: category subtotal does not reconcile")
        for i, metric in enumerate(metrics):
            parts = [v[i] for k, v in observed.items() if k != "Grand Total"]
            total = observed["Grand Total"][i]
            if total is not None and all(v is not None for v in parts) and abs(sum(parts) - total) > 0.02:
                raise ValueError(f"NDMA national total does not reconcile: {metric}")
        for region, cells in observed.items():
            for metric, value in zip(metrics, cells):
                output.append(dict(observation_id=f"{page['report_id']}:{page['page']}:{region}:{metric}",
                    report_id=page["report_id"], source="ndma", location_name="Pakistan" if region == "Grand Total" else region,
                    admin_level="country" if region == "Grand Total" else "province_or_territory",
                    period_start=start, period_end=end, period_type="cumulative", metric=metric,
                    value=value, unit="km" if metric.endswith("_km") else "count",
                    source_page=row_pages[region], source_url=page["source_url"], raw_sha256=page["raw_sha256"],
                    measurement_note="Cumulative monsoon report snapshot. Do not sum across reports, "
                    "country/province levels, or demographic subtotals and totals. '-' is missing, not zero."))
    return output


def fetch_ndma(store, max_pages, pdf_limit):
    found = {}
    complete = False
    for page in range(1, max_pages + 1):
        url = "https://ndma.gov.pk/sitreps?" + urlencode(dict(cat_id=3, page=page))
        data, meta = store.get(url, "ndma")
        links = ndma_links(data, url)
        for link, title in links:
            found[link] = (title, meta)
        soup = BeautifulSoup(data, "html.parser")
        if not any(a.get_text(strip=True).lower() == "next" for a in soup.find_all("a", href=True)):
            complete = True
            break
    reports, assets, pages, impacts, errors = [], [], [], [], []
    # The listing contains dates in human text. Sort by parsed date, not hashed PDF filename.
    def published(title):
        m = re.search(r"\b(\d{2} [A-Za-z]{3} \d{4})\b", title)
        return datetime.strptime(m[1], "%d %b %Y").date().isoformat() if m else None
    selected = sorted(found.items(), key=lambda item: published(item[1][0]) or "", reverse=True)
    for i, (link, (title, meta)) in enumerate(selected):
        rid = "ndma:" + digest(link.encode())[:20]
        reports.append(dict(report_id=rid, source="ndma", title=title,
                            published_at=published(title), glide=None, description=None, source_notes=None,
                            source_url=link, raw_sha256=meta["sha256"], retrieved_at=meta["retrieved_at"],
                            record_note="Situation report, not an independent event. Cumulative totals require review."))
        asset = dict(asset_id=digest(link.encode()), report_id=rid, source="ndma", kind="pdf",
                     source_url=link, download_status="catalogued", raw_path=None, raw_sha256=None)
        if i < pdf_limit:
            try:
                data, pdfmeta = store.get(link, "ndma")
                if not data.startswith(b"%PDF"):
                    raise ValueError("Expected PDF bytes")
                asset.update(download_status="downloaded", raw_path=pdfmeta["path"], raw_sha256=pdfmeta["sha256"])
                extracted = pdf_pages(data, rid, pdfmeta)
                pages.extend(extracted)
                impacts.extend(ndma_impacts(extracted))
            except Exception as exc:
                asset["download_status"] = "extraction_failed" if asset["raw_path"] else "download_failed"
                errors.append(dict(source="ndma", url=link, error=str(exc)))
        assets.append(asset)
    print(f"NDMA: {len(reports)} reports catalogued, {len(pages)} pages extracted", flush=True)
    return dict(reports=reports, assets=assets, pages=pages, impacts=impacts, errors=errors,
                coverage=dict(listing="current monsoon category, not all historical disasters",
                              listing_complete=complete, max_pages=max_pages, pdf_limit=pdf_limit))


def emdat_date(row, prefix):
    parts = [row.get(prefix + " " + x) for x in ("Year", "Month", "Day")]
    if not parts[0]:
        return None, "unknown"
    if not parts[1]:
        return None, "year"
    if not parts[2]:
        return None, "month"
    return date(*(int(x) for x in parts)).isoformat(), "day"


def import_emdat(path, store):
    import openpyxl
    data = Path(path).read_bytes()
    _, meta = store.record(Path(path).resolve().as_uri(), data, "emdat", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    w = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    required = {"DisNo.", "ISO", "Disaster Type", "Start Year"}
    recognized = False
    for sheet in w:
        header = None
        for values in sheet.values:
            if header is None:
                if required.issubset(set(values)):
                    header = list(values)
                    recognized = True
                continue
            r = dict(zip(header, values))
            if r.get("ISO") != "PAK" or r.get("Disaster Type") not in EMDAT_HAZARDS:
                continue
            if not r.get("DisNo."):
                raise ValueError("EM-DAT Pakistan row lacks DisNo.")
            start, precision = emdat_date(r, "Start")
            end, _ = emdat_date(r, "End")
            damage = number(r.get("Total Damage ('000 US$)"))
            rows.append(dict(record_id=f"emdat:{r['DisNo.']}", source="emdat", source_event_id=str(r["DisNo."]),
                             country_iso3="PAK", hazard=norm(r["Disaster Type"]).replace(" ", "_"),
                             subtype=r.get("Disaster Subtype"), title=r.get("Event Name") or r["Disaster Type"],
                             start_date=start, end_date=end, date_precision=precision,
                             start_year=int(r["Start Year"]) if r.get("Start Year") else None,
                             glide=None, alert_level=None, deaths=number(r.get("Total Deaths")),
                             affected=number(r.get("Total Affected")), damage_usd=damage * 1000 if damage is not None else None,
                             source_url="https://public.emdat.be/", source_modified=str(r.get("Last Update") or ""),
                             episode_id=None, raw_sha256=meta["sha256"], retrieved_at=meta["retrieved_at"],
                             coverage_note="EM-DAT inclusion thresholds apply. Restricted reuse terms; local-only output. "
                             "Original date components and location text retained in raw export."))
    w.close()
    if not recognized:
        raise ValueError("No EM-DAT header found; expected " + ", ".join(sorted(required)))
    return rows


# Explicit schemas keep all-null numeric columns numeric, including empty source runs.
SCHEMAS = {
    "events": "record_id source source_event_id country_iso3 hazard subtype title start_date end_date date_precision start_year:BIGINT glide alert_level deaths:DOUBLE affected:DOUBLE damage_usd:DOUBLE source_url source_modified episode_id raw_sha256 retrieved_at coverage_note",
    "reports": "report_id source title published_at glide description source_notes source_url raw_sha256 retrieved_at record_note",
    "assets": "asset_id report_id source kind source_url download_status raw_path raw_sha256",
    "exposures": "observation_id report_id source location_name admin_level province district_key match_status observation_start observation_end area_km2:DOUBLE analyzed_area_km2:DOUBLE analyzed_fraction:DOUBLE flooded_area_km2:DOUBLE population_total:DOUBLE population_analyzed:DOUBLE population_exposed:DOUBLE exposed_fraction:DOUBLE source_url source_sheet source_row:BIGINT raw_sha256 measurement_note",
    "pages": "report_id page:BIGINT text raw_sha256 source_url",
    "impacts": "observation_id report_id source location_name admin_level period_start period_end period_type metric value:DOUBLE unit source_page:BIGINT source_url raw_sha256 measurement_note",
}
KEYS = dict(events="record_id", reports="report_id", assets="asset_id", exposures="observation_id", impacts="observation_id")


def dedupe_events(rows):
    latest = {}
    for r in rows:
        key = r["record_id"]
        # Newest source revision wins; never choose maximum casualty count.
        rank = (r.get("source_modified") or "", r.get("retrieved_at") or "", int(r.get("episode_id") or 0))
        if key not in latest or rank > latest[key][0]:
            latest[key] = (rank, r)
    return [latest[k][1] for k in sorted(latest)]


def build(root, out):
    # Concurrent collectors may finish together. Serialize output replacement.
    import fcntl
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    with (out / ".build.lock").open("w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        return _build(root, out)


def _build(root, out):
    import duckdb
    tables = {k: [] for k in SCHEMAS}
    statuses = []
    runs = [json.loads(path.read_text()) for path in (Path(root) / "runs").glob("*.json")]
    for run in sorted(runs, key=lambda r: r.get("completed_at") or ""):
        statuses.append({k: run.get(k) for k in ("source", "completed_at", "errors", "coverage")})
        for name in tables:
            tables[name].extend(run.get(name, []))
    tables["events"] = dedupe_events(tables["events"])
    for name in ("reports", "assets", "exposures", "impacts"):
        by_key = {r[KEYS[name]]: r for r in tables[name]}
        tables[name] = [by_key[k] for k in sorted(by_key)]
    tables["pages"] = list({(r["report_id"], r["page"]): r for r in tables["pages"]}.values())
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    catalog = []
    for name, schema in SCHEMAS.items():
        columns = [(x.split(":")[0], x.split(":")[1] if ":" in x else "VARCHAR") for x in schema.split()]
        names = [x[0] for x in columns]
        con.execute(f'CREATE TABLE "{name}" (' + ",".join(f'"{k}" {t}' for k, t in columns) + ")")
        if tables[name]:
            con.executemany(f'INSERT INTO "{name}" VALUES (' + ",".join("?" for _ in names) + ")",
                            [[r.get(k) for k in names] for r in tables[name]])
        if name == "events":
            invalid = con.sql("SELECT count(*) FROM events WHERE country_iso3 <> 'PAK' OR (start_date IS NOT NULL AND end_date IS NOT NULL AND start_date > end_date)").fetchone()[0]
            if invalid:
                raise ValueError(f"{invalid} invalid event rows")
        for ext, options in [("parquet", "FORMAT PARQUET, COMPRESSION ZSTD"), ("csv", "FORMAT CSV, HEADER TRUE")]:
            dest = out / f"climate_{name}.{ext}"
            temp = dest.with_suffix(dest.suffix + ".tmp")
            escaped = str(temp).replace("'", "''")
            con.execute(f"COPY {name} TO '{escaped}' ({options})")
            temp.replace(dest)
        catalog.append(dict(name=f"climate_{name}", file=f"climate_{name}.parquet", rows=len(tables[name]),
                            columns=[dict(name=k, type=t) for k, t in columns]))
    con.close()
    unmatched = [dict(location_name=r["location_name"], province=r["province"], report_id=r["report_id"])
                 for r in tables["exposures"] if r["admin_level"] == "district" and not r["district_key"]]
    anomalies = [dict(observation_id=r["observation_id"], location_name=r["location_name"],
                      issue="Source fraction exceeds 1; retained without clipping")
                 for r in tables["exposures"]
                 if any(r.get(k) is not None and r[k] > 1 for k in ("analyzed_fraction", "exposed_fraction"))]
    quality = dict(generated_at=now(), row_counts={k: len(v) for k, v in tables.items()},
                   sources=statuses, unmatched_districts=unmatched,
                   source_anomalies=anomalies,
                   empty_pdf_pages=sum(not r["text"].strip() for r in tables["pages"]),
                   pages_needing_ocr_or_visual_review=[dict(report_id=r["report_id"], page=r["page"])
                       for r in tables["pages"] if not r["text"].strip()],
                   ndma_reports_without_structured_impacts=sorted(
                       {r["report_id"] for r in tables["pages"] if "ndma.gov.pk" in r["source_url"]}
                       - {r["report_id"] for r in tables["impacts"]}),
                   publication="Local warehouse only; source-specific reuse terms not replaced by repo licence.")
    write_json(out / "catalog.json", dict(name="Pakistan weather-related events", generated=now(), tables=catalog,
        notes="Source-specific event records may describe the same disaster. Do not sum across sources. "
              "Reports are not events. Exposure is not confirmed impact. Null is not zero."))
    write_json(out / "quality_report.json", quality)
    print(json.dumps(quality["row_counts"], indent=2))
    return quality


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--raw", type=Path, default=WORKSPACE / "raw_data/climate_events")
    ap.add_argument("--out", type=Path, default=WORKSPACE / "data_darbar_warehouse/climate_events")
    ap.add_argument("--refresh", action="store_true", help="Download new immutable snapshots")
    ap.add_argument("--offline", action="store_true", help="Only use cached source files")
    subs = ap.add_subparsers(dest="command", required=True)
    p = subs.add_parser("gdacs", help="Fetch Pakistan floods, cyclones, droughts and wildfires")
    p.add_argument("--start", type=date.fromisoformat, default=date(2000, 1, 1))
    p.add_argument("--end", type=date.fromisoformat, default=date.today())
    p = subs.add_parser("unosat", help="Download product metadata, PDFs and exposure XLSX")
    p.add_argument("--ids", nargs="+", type=int, default=[3343])
    p.add_argument("--geometry", action="store_true", help="Also download linked GIS files")
    p = subs.add_parser("ndma", help="Catalogue current monsoon reports and extract PDF pages")
    p.add_argument("--max-pages", type=int, default=20)
    p.add_argument("--pdf-limit", type=int, default=3)
    p = subs.add_parser("emdat", help="Import an authorised local EM-DAT Excel export")
    p.add_argument("file", type=Path)
    subs.add_parser("build", help="Rebuild local CSV/Parquet from extracted runs, without network")
    args = ap.parse_args()
    if args.command == "build":
        build(args.raw, args.out)
        return
    store = Store(args.raw, args.refresh, args.offline)
    if args.command == "gdacs":
        if args.start > args.end:
            ap.error("--start must precede --end")
        result = dict(events=fetch_gdacs(store, args.start, args.end),
                      coverage=dict(requested_start=str(args.start), requested_end=str(args.end),
                                    note="GDACS coverage varies by hazard/year; no heatwave series."))
        run_name = f"gdacs_{args.start}_{args.end}"
    elif args.command == "unosat":
        known = set(json.loads((REPO / "app/data/districts.json").read_text()))
        result = fetch_unosat(store, args.ids, known, args.geometry)
        run_name = "unosat_" + "_".join(str(x) for x in sorted(args.ids))
    elif args.command == "ndma":
        if args.max_pages < 1 or args.pdf_limit < 0:
            ap.error("--max-pages must be positive and --pdf-limit nonnegative")
        result = fetch_ndma(store, args.max_pages, args.pdf_limit)
        run_name = "ndma"
    else:
        result = dict(events=import_emdat(args.file, store), coverage=dict(file=args.file.name))
        run_name = "emdat"
    result.update(source=args.command, completed_at=now())
    write_json(args.raw / "runs" / (run_name + ".json"), result)
    build(args.raw, args.out)
    if result.get("errors"):
        raise SystemExit("Partial extraction: inspect quality_report.json; cached successful files are reusable.")


if __name__ == "__main__":
    main()
